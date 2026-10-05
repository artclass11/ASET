from csv import DictReader
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def test_official_source_registry_is_traceable() -> None:
    registry = ROOT / "data" / "official_source_registry.csv"
    with registry.open(newline="", encoding="utf-8") as handle:
        rows = list(DictReader(handle))
    assert len(rows) >= 10
    assert all(row["source_ref"] and row["publisher"] and row["official_url"] for row in rows)
    assert all(row["source_tier"] in {"1", "2", "3"} for row in rows)


def test_workbook_is_original_blank_formula_driven_template() -> None:
    workbook_path = ROOT / "deliverables" / "ASET_Professional_Global_Equity_Research_Template.xlsx"
    workbook = load_workbook(workbook_path, data_only=False)
    assert workbook.sheetnames == ["Cover", "Screen_Input", "Screen_Output", "Research_Card", "Source_Manifest"]
    assert workbook["Cover"]["B5"].value.startswith("TEMPLATE")
    assert workbook["Screen_Input"]["A5"].value is None
    assert "Screen_Input" in workbook["Screen_Output"]["A5"].value
    assert workbook["Screen_Output"]["J5"].value.startswith("=IF(")
    assert workbook["Cover"]["A1"].fill.fgColor.rgb.endswith("0B0B0D")


def test_pdf_is_readable_and_contains_notice() -> None:
    pdf_path = ROOT / "deliverables" / "ASET_Professional_Global_Equity_Research_Framework.pdf"
    assert pdf_path.stat().st_size > 2000
    import subprocess

    text = subprocess.check_output(["pdftotext", str(pdf_path), "-"], text=True)
    assert "OFFICIAL-SOURCE POLICY" in text
    assert "IMPORTANT NOTICE" in text
    assert "no market figures" in text
